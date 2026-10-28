from pathlib import Path
from typing import Callable

from openpyxl import load_workbook

from .models import WorkbookSheetSummary, WorkbookSummary


ProgressCallback = Callable[[int, str], None]
SUPPORTED_FORMATS = {".xls", ".xlsx", ".xlsm"}


def inspect_workbook(path: str | Path, progress: ProgressCallback | None = None) -> WorkbookSummary:
    workbook_path = Path(path)
    if workbook_path.suffix.lower() not in SUPPORTED_FORMATS:
        supported = ", ".join(sorted(SUPPORTED_FORMATS))
        raise ValueError(f"Unsupported workbook format. Choose one of: {supported}.")
    if not workbook_path.is_file():
        raise FileNotFoundError(f"Workbook not found: {workbook_path}")

    if workbook_path.suffix.lower() == ".xls":
        return _inspect_xls(workbook_path, progress)
    return _inspect_openpyxl(workbook_path, progress)


def _inspect_openpyxl(path: Path, progress: ProgressCallback | None) -> WorkbookSummary:
    formula_book = load_workbook(path, read_only=True, data_only=False)
    value_book = load_workbook(path, read_only=True, data_only=True)
    summaries: list[WorkbookSheetSummary] = []
    workbook_errors: list[str] = []
    sheet_count = len(formula_book.worksheets)

    try:
        for sheet_index, formula_sheet in enumerate(formula_book.worksheets, start=1):
            value_sheet = value_book[formula_sheet.title]
            populated_rows = 0
            formula_count = 0
            sheet_errors: list[str] = []
            rows = max(formula_sheet.max_row or 0, value_sheet.max_row or 0)
            columns = max(formula_sheet.max_column or 0, value_sheet.max_column or 0)

            formula_rows = formula_sheet.iter_rows(max_row=rows, max_col=columns)
            value_rows = value_sheet.iter_rows(max_row=rows, max_col=columns)
            for row_number, (formula_row, value_row) in enumerate(zip(formula_rows, value_rows), start=1):
                row_has_data = False
                for formula_cell, value_cell in zip(formula_row, value_row):
                    if formula_cell.value is not None:
                        row_has_data = True
                    if formula_cell.data_type == "f":
                        formula_count += 1
                    if value_cell.data_type == "e":
                        message = f"{formula_sheet.title}!{value_cell.coordinate}: {value_cell.value}"
                        if len(sheet_errors) < 100:
                            sheet_errors.append(message)
                if row_has_data:
                    populated_rows += 1

            summaries.append(
                WorkbookSheetSummary(
                    name=formula_sheet.title,
                    rows=rows,
                    columns=columns,
                    populated_rows=populated_rows,
                    formula_count=formula_count,
                    errors=tuple(sheet_errors),
                )
            )
            workbook_errors.extend(sheet_errors)
            if progress:
                percent = int(sheet_index * 100 / max(sheet_count, 1))
                progress(percent, f"Read sheet {sheet_index} of {sheet_count}: {formula_sheet.title}")
    finally:
        formula_book.close()
        value_book.close()

    return WorkbookSummary(str(path), tuple(summaries), tuple(workbook_errors))


def _inspect_xls(path: Path, progress: ProgressCallback | None) -> WorkbookSummary:
    try:
        import xlrd
    except ImportError as error:
        raise RuntimeError("Reading .xls files requires xlrd. Install the project requirements first.") from error

    workbook = xlrd.open_workbook(path, on_demand=True)
    summaries: list[WorkbookSheetSummary] = []
    workbook_errors: list[str] = []
    try:
        for sheet_index, sheet in enumerate(workbook.sheets(), start=1):
            populated_rows = 0
            sheet_errors: list[str] = []
            for row_index in range(sheet.nrows):
                row_has_data = False
                for column_index in range(sheet.ncols):
                    cell = sheet.cell(row_index, column_index)
                    if cell.ctype != xlrd.XL_CELL_EMPTY:
                        row_has_data = True
                    if cell.ctype == xlrd.XL_CELL_ERROR:
                        address = f"{sheet.name}!R{row_index + 1}C{column_index + 1}"
                        message = f"{address}: {xlrd.error_text_from_code.get(cell.value, 'Excel error')}"
                        if len(sheet_errors) < 100:
                            sheet_errors.append(message)
                if row_has_data:
                    populated_rows += 1

            summaries.append(
                WorkbookSheetSummary(
                    name=sheet.name,
                    rows=sheet.nrows,
                    columns=sheet.ncols,
                    populated_rows=populated_rows,
                    formula_count=0,
                    errors=tuple(sheet_errors),
                )
            )
            workbook_errors.extend(sheet_errors)
            if progress:
                percent = int(sheet_index * 100 / max(workbook.nsheets, 1))
                progress(percent, f"Read sheet {sheet_index} of {workbook.nsheets}: {sheet.name}")
    finally:
        workbook.release_resources()

    return WorkbookSummary(str(path), tuple(summaries), tuple(workbook_errors))