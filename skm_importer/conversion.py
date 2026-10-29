from collections import defaultdict
import logging
from pathlib import Path
from typing import Any

from openpyxl import load_workbook

from .rpt_schema import SECTION_DEFAULTS, SECTION_FIELDS


def build_records(
    workbook_path: str | Path,
    logger: logging.Logger | None = None,
) -> tuple[dict[str, list[tuple[str, ...]]], tuple[str, ...]]:
    path = Path(workbook_path)
    log = logger or logging.getLogger(__name__)
    log.info("Opening workbook for SKM conversion: %s", path)
    if path.suffix.lower() not in {".xlsx", ".xlsm"}:
        log.error("Unsupported input extension: %s", path.suffix)
        raise ValueError("Generating SKM reports requires an .xlsx or .xlsm workbook.")
    if not path.is_file():
        log.error("Input workbook does not exist: %s", path)
        raise FileNotFoundError(f"Workbook not found: {path}")

    workbook = load_workbook(path, read_only=True, data_only=True)
    try:
        log.info("Workbook opened; sheets: %s", workbook.sheetnames)
        if "FIELD CABLEOCP" not in workbook.sheetnames:
            log.error("Required sheet FIELD CABLEOCP was not found")
            raise ValueError("The workbook is missing the FIELD CABLEOCP sheet.")

        records: dict[str, list[tuple[str, ...]]] = defaultdict(list)
        warnings: list[str] = []
        log.info("Converting transformer survey rows")
        transformer_sides = _transformers(workbook, records, warnings, log)
        buses: dict[str, dict[str, str]] = {}
        log.info("Converting bus duct survey rows")
        _bus_ducts(workbook, buses, warnings, log)
        utilities: dict[str, tuple[str, str]] = {}
        cable_names: set[str] = set()
        device_names: set[str] = set()
        cable_indexes: dict[str, int] = {}
        incoming_cables: dict[str, str] = {}
        sheet = workbook["FIELD CABLEOCP"]

        for row_number, values in enumerate(sheet.iter_rows(min_row=6, values_only=True), start=6):
            source_type, source_label = _text(_cell(values, 1)), _text(_cell(values, 2))
            load_type, load_label = _text(_cell(values, 3)), _text(_cell(values, 4))
            if not any((source_type, source_label, load_type, load_label)):
                continue
            row_context = f"FIELD CABLEOCP row {row_number}"
            log.debug(
                "%s inputs: %s",
                row_context,
                {
                    "source_type": source_type,
                    "source_label": source_label,
                    "load_type": load_type,
                    "load_label": load_label,
                    "voltage": _text(_cell(values, 5)),
                    "conductors": _text(_cell(values, 9)),
                    "cable_size": _text(_cell(values, 10)),
                    "length": _text(_cell(values, 11)),
                    "conductor_material": _text(_cell(values, 12)),
                    "cable_type": _text(_cell(values, 13)),
                    "raceway": _text(_cell(values, 14)),
                    "feeder_configuration": _text(_cell(values, 15)),
                    "protection_manufacturer": _text(_cell(values, 19)),
                    "protection_model": _text(_cell(values, 20)),
                    "interrupting_rating": _text(_cell(values, 21)),
                    "sensor_or_plug_rating": _text(_cell(values, 22)),
                    "breaker_settings": {
                        column: _text(_cell(values, index))
                        for column, index in zip("XYZ", (23, 24, 25))
                    } | {
                        column: _text(_cell(values, index))
                        for column, index in zip(("AA", "AB", "AC", "AD", "AE"), range(26, 31))
                    },
                },
            )

            protection_values = tuple(_text(_cell(values, column)) for column in (19, 20, 21, 22))
            has_protection = any(protection_values)
            settings_values = {
                column: _text(_cell(values, index))
                for column, index in zip(("X", "Y", "Z", "AA", "AB", "AC", "AD", "AE"), range(23, 31))
                if _text(_cell(values, index))
            }
            is_main_protection = has_protection and (
                load_type in {"CB MAIN", "FUSE MAIN"} or not load_type
            )
            if source_type and source_label and is_main_protection:
                source_ref = _endpoint_ref(source_type, source_label, _text(_cell(values, 5)), transformer_sides)
                if source_ref is None:
                    _warn(warnings, f"{row_context}: main protective device has an unresolved source endpoint.", log)
                    continue
                incoming_cable_name = incoming_cables.get(source_ref)
                if incoming_cable_name is None:
                    _warn(warnings, f"{row_context}: no incoming feeder was found for the main protective device.", log)
                    continue

                cable_index = cable_indexes[incoming_cable_name]
                cable_fields = SECTION_FIELDS["Cable"]
                incoming_record = dict(zip(cable_fields, records["Cable"][cable_index]))
                if incoming_record["ConnectedComponent2"] != source_ref:
                    _warn(
                        warnings,
                        f"{row_context}: incoming feeder {incoming_cable_name} already terminates at {incoming_record['ConnectedComponent2']}; main device was not inserted.",
                        log,
                    )
                    continue

                equipment_label = f"{source_type} {source_label}"
                model = _text(_cell(values, 20))
                device_name = _device_name(load_type, equipment_label, model)
                if device_name in device_names:
                    device_name = f"{device_name} {row_number}"
                    _warn(warnings, f"{row_context}: duplicate main-device name; generated {device_name}.", log)
                device_names.add(device_name)
                if settings_values:
                    _warn(
                        warnings,
                        f"{row_context}: breaker settings {settings_values} were not translated because the SKM setting syntax is model-specific.",
                        log,
                    )

                voltage = buses.get(_node_name(source_type, source_label), {}).get("SystemNominalVoltage", "")
                if not voltage:
                    voltage = _text(_cell(values, 5))
                protection_code = _protection_code(device_name, model)
                if not protection_code:
                    _warn(warnings, f"{row_context}: protective-device type is unknown; left its SKM type blank.", log)
                incoming_record["ConnectedComponent2"] = f"{device_name}:1"
                records["Cable"][cable_index] = tuple(incoming_record[field] for field in cable_fields)
                log.debug(
                    "%s inserted %s between feeder %s and bus %s:1",
                    row_context,
                    device_name,
                    incoming_cable_name,
                    source_ref.removesuffix(":1"),
                )
                _append_record(
                    records,
                    "Protective Device",
                    {
                        "ComponentName": device_name,
                        "ComponentType": "112",
                        "ConnectedComponent1": f"{incoming_cable_name}:2",
                        "ConnectedComponent2": source_ref,
                        "SystemNominalVoltage": voltage,
                        "ProtectionType": protection_code,
                        "Manufacturer": _text(_cell(values, 19)),
                        "Type": model,
                        "FrameVoltage": voltage,
                        "InterruptingRating": _text(_cell(values, 21)),
                        "Sensor/Trip": _text(_cell(values, 22)),
                        "Plug": _text(_cell(values, 22)),
                    },
                    logger=log,
                    context=row_context,
                )
                continue

            if not all((source_type, source_label, load_type, load_label)):
                _warn(warnings, f"{row_context}: skipped because a source or load type/label is missing.", log)
                continue

            voltage = _text(_cell(values, 5))
            source_ref = _endpoint_ref(source_type, source_label, voltage, transformer_sides)
            load_ref = _endpoint_ref(load_type, load_label, voltage, transformer_sides)
            if source_ref is None or load_ref is None:
                _warn(warnings, f"{row_context}: skipped because a transformer endpoint is ambiguous or missing from XFMR.", log)
                continue
            log.debug("%s endpoint resolution: source=%s load=%s", row_context, source_ref, load_ref)

            for component_type, label in ((source_type, source_label), (load_type, load_label)):
                if component_type == "TRANS":
                    continue
                bus_name = _node_name(component_type, label)
                bus = buses.setdefault(bus_name, {"ComponentName": bus_name, "ComponentType": "10"})
                if voltage and not bus.get("SystemNominalVoltage"):
                    bus["SystemNominalVoltage"] = voltage
                elif voltage and bus.get("SystemNominalVoltage") != voltage:
                    _warn(
                        warnings,
                        f"{row_context}: {bus_name} has conflicting nominal voltages; kept {bus['SystemNominalVoltage']}.",
                        log,
                    )
                if component_type == "UTILITY":
                    utilities[bus_name] = (f"UTILITY-{label}", voltage)

            cable_name = f"FEED {load_type} {load_label}"
            if cable_name in cable_names:
                _warn(warnings, f"{row_context}: duplicate feeder name {cable_name}; skipped.", log)
                continue
            cable_names.add(cable_name)

            device_name = _device_name(load_type, load_label, _text(_cell(values, 20)))
            if has_protection:
                if device_name in device_names:
                    device_name = f"{device_name} {row_number}"
                    _warn(warnings, f"{row_context}: duplicate protection name; generated {device_name}.", log)
                device_names.add(device_name)
                protection_code = _protection_code(device_name, _text(_cell(values, 20)))
                if not protection_code:
                    _warn(warnings, f"{row_context}: protective-device type is unknown; left its SKM type blank.", log)
                _append_record(
                    records,
                    "Protective Device",
                    {
                        "ComponentName": device_name,
                        "ComponentType": "112",
                        "ConnectedComponent1": f"{cable_name}:2",
                        "ConnectedComponent2": load_ref,
                        "SystemNominalVoltage": voltage,
                        "ProtectionType": protection_code,
                        "Manufacturer": _text(_cell(values, 19)),
                        "Type": _text(_cell(values, 20)),
                        "FrameVoltage": voltage,
                        "InterruptingRating": _text(_cell(values, 21)),
                        "Sensor/Trip": _text(_cell(values, 22)),
                        "Plug": _text(_cell(values, 22)),
                    },
                    logger=log,
                    context=row_context,
                )

            _append_record(
                records,
                "Cable",
                {
                    "ComponentName": cable_name,
                    "ComponentType": "20",
                    "ConnectedComponent1": source_ref,
                    "ConnectedComponent2": f"{device_name}:1" if has_protection else load_ref,
                    "SystemNominalVoltage": voltage,
                    "CableSize": _text(_cell(values, 10)),
                    "Length": _text(_cell(values, 11)),
                    "ConductorType": _conductor_type(_cell(values, 12)),
                    "InsulationClass": _text(_cell(values, 13)),
                    "RacewayType": _text(_cell(values, 14)),
                    "Conductor Desc": _text(_cell(values, 15)),
                    "#/Conductors": _conductor_count(_cell(values, 9)),
                },
                logger=log,
                context=row_context,
            )
            cable_indexes[cable_name] = len(records["Cable"]) - 1
            incoming_cables[load_ref] = cable_name
            if not _text(_cell(values, 10)) or not _text(_cell(values, 11)):
                _warn(warnings, f"{row_context}: {cable_name} is missing cable size or length.", log)

        for bus in buses.values():
            _append_record(records, "Bus", bus, logger=log, context="FIELD CABLEOCP bus aggregation")
        for bus_name, (utility_name, voltage) in utilities.items():
            _append_record(
                records,
                "Utility",
                {
                    "ComponentName": utility_name,
                    "ComponentType": "70",
                    "ConnectedComponent1": f"{bus_name}:1",
                    "SystemNominalVoltage": voltage,
                },
                logger=log,
                context="UTILITY endpoint aggregation",
            )

        log.info("Conversion completed; records by section: %s", {name: len(rows) for name, rows in records.items()})
        log.info("Conversion completed with %d warning(s)", len(warnings))
        return dict(records), tuple(warnings)
    finally:
        workbook.close()


def _transformers(
    workbook: Any,
    records: dict[str, list[tuple[str, ...]]],
    warnings: list[str],
    logger: logging.Logger,
) -> dict[str, tuple[str, str, str, str]]:
    sides: dict[str, tuple[str, str, str, str]] = {}
    unmapped_construction_types = 0
    if "XFMR" not in workbook.sheetnames:
        logger.info("XFMR sheet not present; no transformer records to convert")
        return sides

    for row_number, values in enumerate(workbook["XFMR"].iter_rows(min_row=6, values_only=True), start=6):
        name = _text(_cell(values, 1))
        if not name:
            continue
        context = f"XFMR row {row_number}"
        component_name = name if name.upper().startswith("TRANS ") else f"TRANS {name}"
        label = component_name[6:].strip()
        primary_bus, secondary_bus = f"{label} PRIM BUS", label
        primary_voltage, secondary_voltage = _text(_cell(values, 2)), _text(_cell(values, 3))
        if not primary_voltage or not secondary_voltage:
            _warn(warnings, f"{context}: {component_name} is missing a primary or secondary voltage.", logger)
        if _text(_cell(values, 6)):
            unmapped_construction_types += 1
        sides[component_name.casefold()] = (primary_bus, secondary_bus, primary_voltage, secondary_voltage)
        resistance, reactance = _impedance_components(_text(_cell(values, 5)), _text(_cell(values, 8)))
        primary_connection, secondary_connection = _connection_codes(_text(_cell(values, 7)))
        _append_record(
            records,
            "2-Winding Transformer",
            {
                "ComponentName": component_name,
                "ComponentType": "30",
                "ConnectedComponent1": f"{primary_bus}:1",
                "ConnectedComponent2": f"{secondary_bus}:1",
                "SystemNominalVoltage": primary_voltage,
                "SystemNominalVoltageSecondary": secondary_voltage,
                "Nominal kVA": _text(_cell(values, 4)),
                "Pri Connection": primary_connection,
                "Sec Connection": secondary_connection,
                "Rpos": resistance,
                "Xpos": reactance,
            },
            logger=logger,
            context=context,
        )
        _append_record(records, "Bus", {"ComponentName": primary_bus, "ComponentType": "10", "SystemNominalVoltage": primary_voltage}, logger=logger, context=context)
        _append_record(records, "Bus", {"ComponentName": secondary_bus, "ComponentType": "10", "SystemNominalVoltage": secondary_voltage}, logger=logger, context=context)
    if unmapped_construction_types:
        _warn(
            warnings,
            f"XFMR: survey construction type was not mapped for {unmapped_construction_types} records; SKM transformer type fields were left blank.",
            logger,
        )
    return sides


def _bus_ducts(
    workbook: Any,
    buses: dict[str, dict[str, str]],
    warnings: list[str],
    logger: logging.Logger,
) -> None:
    if "BUS DUCT" not in workbook.sheetnames:
        logger.info("BUS DUCT sheet not present; no bus duct records to convert")
        return
    for row_number, values in enumerate(workbook["BUS DUCT"].iter_rows(min_row=6, values_only=True), start=6):
        name = _text(_cell(values, 1))
        if not name:
            continue
        context = f"BUS DUCT row {row_number}"
        component_name = name if name.upper().startswith("BUS RISER ") else f"BUS RISER {name}"
        bus = buses.setdefault(component_name, {"ComponentName": component_name, "ComponentType": "10"})
        bus.update(
            {
                "ContinuousRating": _text(_cell(values, 3)),
                "Manufacturer": _text(_cell(values, 5)),
                "Description": _text(_cell(values, 2)),
            }
        )
        logger.debug("%s mapped to Bus record %s: %s", context, component_name, bus)
        if not bus["ContinuousRating"]:
            _warn(warnings, f"{context}: {component_name} is missing ampacity.", logger)


def _endpoint_ref(
    component_type: str,
    label: str,
    voltage: str,
    transformer_sides: dict[str, tuple[str, str, str, str]],
) -> str | None:
    if component_type != "TRANS":
        return f"{_node_name(component_type, label)}:1"
    sides = transformer_sides.get(f"TRANS {label}".casefold())
    if sides is None:
        return None
    primary_bus, secondary_bus, primary_voltage, secondary_voltage = sides
    if voltage and voltage == primary_voltage:
        return f"{primary_bus}:1"
    if voltage and voltage == secondary_voltage:
        return f"{secondary_bus}:1"
    return None


def _node_name(component_type: str, label: str) -> str:
    return f"UTILITY {label}" if component_type == "UTILITY" else f"{component_type} {label}"


def _device_name(load_type: str, load_label: str, model: str) -> str:
    model_name = model.upper()
    fuse_models = {
        "A4BY", "A4J", "A6D", "AJT", "CLF", "ECSR", "FLSRID", "FRN-R", "FRS-R",
        "JKS", "LPJ_SP", "LPN-RK", "LPS-RK", "TR", "TRS",
    }
    words = load_type.strip().split(maxsplit=1)
    declared_prefix = words[0].upper() if words else ""
    category = words[1] if declared_prefix in {"CB", "FUSE"} and len(words) > 1 else load_type
    prefix = "FUSE" if declared_prefix == "FUSE" or model_name in fuse_models else "CB"
    parts = [prefix]
    if category:
        parts.append(category)
    if load_label:
        parts.append(load_label)
    return " ".join(parts)


def _protection_code(device_name: str, model: str) -> str:
    if device_name.startswith("FUSE "):
        return "11041"
    if model.upper() in {"TP", "THP", "SKPB"}:
        return "11001"
    if model.upper() in {"EHB", "LA", "NF", "QB", "QP", "TED", "TEY", "THQB", "TQD"}:
        return "11003"
    return ""


def _append_record(
    records: dict[str, list[tuple[str, ...]]],
    section: str,
    values: dict[str, str],
    logger: logging.Logger | None = None,
    context: str = "",
) -> None:
    log = logger or logging.getLogger(__name__)
    log.debug("Building %s record%s; mapped fields: %s", section, f" from {context}" if context else "", values)
    record = {field: "" for field in SECTION_FIELDS[section]}
    record.update(SECTION_DEFAULTS.get(section, {}))
    for field, value in values.items():
        if field not in record:
            raise ValueError(f"Unknown {section} field: {field}")
        record[field] = value
    records.setdefault(section, []).append(tuple(record[field] for field in SECTION_FIELDS[section]))
    log.debug("Built %s record %s with %d fields", section, record.get("ComponentName", "<unnamed>"), len(record))


def _warn(warnings: list[str], message: str, logger: logging.Logger) -> None:
    warnings.append(message)
    logger.warning(message)


def _cell(values: tuple[object, ...], index: int) -> object:
    return values[index] if index < len(values) else None


def _text(value: object) -> str:
    if value is None:
        return ""
    if isinstance(value, float) and value.is_integer():
        return str(int(value))
    return str(value).strip()


def _conductor_type(value: object) -> str:
    return {"CU": "Copper", "AL": "Aluminum"}.get(_text(value).upper(), _text(value))


def _conductor_count(value: object) -> str:
    try:
        count = int(float(_text(value)))
    except ValueError:
        return ""
    return str(count - 1) if 1 <= count <= 6 else ""


def _connection_codes(connection: str) -> tuple[str, str]:
    codes = {"DELTA": "0", "WYE": "1", "WYE GND": "2", "WYE-GND": "2"}
    sides = [side.strip() for side in connection.upper().split("-")]
    if len(sides) != 2:
        return "", ""
    return codes.get(sides[0], ""), codes.get(sides[1], "")


def _impedance_components(impedance: str, x_over_r: str) -> tuple[str, str]:
    try:
        magnitude, ratio = float(impedance), float(x_over_r)
    except ValueError:
        return "", ""
    resistance = magnitude / (1 + ratio * ratio) ** 0.5
    reactance = resistance * ratio
    return f"{resistance:.4f}", f"{reactance:.4f}"