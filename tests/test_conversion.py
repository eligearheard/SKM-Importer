import tempfile
import unittest
from pathlib import Path

from openpyxl import Workbook

from skm_importer.conversion import build_records
from skm_importer.pipeline import generate_import_report
from skm_importer.rpt_schema import SECTION_FIELDS


class WorkbookConversionTests(unittest.TestCase):
    def test_builds_skm_records_without_a_reference_report(self) -> None:
        workbook = Workbook()
        cable_sheet = workbook.active
        cable_sheet.title = "FIELD CABLEOCP"
        cable_sheet.append([None] * 32)
        cable_sheet.append([None] * 32)
        cable_sheet.append([None] * 32)
        cable_sheet.append([None] * 32)
        cable_sheet.append([None] * 32)
        cable_sheet.append([1, "TRANS", "TX1", "PNL", "L8C", 208, "VCB", "NO", None, 3, 2, 20, "CU", "XHHW", "Magnetic", "3-1/C+G"])
        main_device = [None] * 32
        main_device[0:4] = [2, "PNL", "L8C", "CB MAIN"]
        main_device[19:23] = ["GE", "TQD", 10, 225]
        cable_sheet.append(main_device)
        cable_sheet.append([3, "UTILITY", "1", "SWBD", "MAIN", 480, None, None, None, 3, 500, 40, "CU", "THHN", "Magnetic", "4-1/C+G"])

        transformer_sheet = workbook.create_sheet("XFMR")
        for _ in range(5):
            transformer_sheet.append([])
        transformer_sheet.append([1, "TRANS TX1", 480, 208, 75, 5.5, "DRY", "DELTA - WYE GND", 3.15])

        bus_duct_sheet = workbook.create_sheet("BUS DUCT")
        for _ in range(5):
            bus_duct_sheet.append([])
        bus_duct_sheet.append([1, "BUS RISER 1", "Spectra", 2000, 100, "GE", "AL"])

        with tempfile.TemporaryDirectory() as temporary_directory:
            path = Path(temporary_directory) / "survey.xlsx"
            output_path = Path(temporary_directory) / "generated.rpt"
            log_path = output_path.with_suffix(".log")
            workbook.save(path)
            records, warnings = build_records(path)
            workbook_summary, report_summary = generate_import_report(path, output_path)
            log_text = log_path.read_text(encoding="utf-8")

        self.assertEqual(len(warnings), 1)
        self.assertIn("construction type was not mapped for 1 records", warnings[0])
        self.assertTrue(output_path.is_file())
        self.assertTrue(log_path.is_file())
        self.assertIn("FIELD CABLEOCP row 6 inputs", log_text)
        self.assertIn("mapped fields", log_text)
        self.assertIn("Generation completed successfully", log_text)
        self.assertEqual(report_summary.errors, warnings)
        self.assertEqual(len(workbook_summary.sheets), 3)
        transformer = dict(zip(SECTION_FIELDS["2-Winding Transformer"], records["2-Winding Transformer"][0]))
        cable = dict(zip(SECTION_FIELDS["Cable"], records["Cable"][0]))
        protection = dict(zip(SECTION_FIELDS["Protective Device"], records["Protective Device"][0]))
        utility = dict(zip(SECTION_FIELDS["Utility"], records["Utility"][0]))
        bus_records = [dict(zip(SECTION_FIELDS["Bus"], record)) for record in records["Bus"]]
        bus_duct = next(bus for bus in bus_records if bus["ComponentName"] == "BUS RISER 1")
        survey_bus = next(bus for bus in bus_records if bus["ComponentName"] == "PNL L8C")
        self.assertEqual(transformer["ComponentName"], "TRANS TX1")
        self.assertEqual(transformer["ConnectedComponent2"], "TX1:1")
        self.assertEqual(transformer["InService"], "1")
        self.assertEqual(transformer["TransformerType"], "0")
        self.assertEqual(transformer["Phase"], "7")
        self.assertEqual(transformer["Size/Do Not Size"], "0")
        self.assertEqual(transformer["SizingCriteria"], "0")
        self.assertEqual(cable["ComponentName"], "FEED PNL L8C")
        self.assertEqual(cable["ConnectedComponent1"], "TX1:1")
        self.assertEqual(cable["ConnectedComponent2"], "CB MAIN PNL L8C:1")
        self.assertEqual(cable["CableSize"], "2")
        self.assertEqual(cable["Length"], "20")
        self.assertEqual(cable["InService"], "1")
        self.assertEqual(cable["Phase"], "7")
        self.assertEqual(cable["Size/Do Not Size"], "0")
        self.assertEqual(cable["RacewayType"], "Magnetic")
        self.assertEqual(protection["ComponentName"], "CB MAIN PNL L8C")
        self.assertEqual(protection["InService"], "1")
        self.assertEqual(protection["ProtectionType"], "11003")
        self.assertEqual(protection["Manufacturer"], "GE")
        self.assertEqual(protection["Type"], "TQD")
        self.assertEqual(protection["InterruptingRating"], "10")
        self.assertEqual(protection["Sensor/Trip"], "225")
        self.assertEqual(protection["Plug"], "225")
        self.assertEqual(protection["Poles"], "0")
        self.assertEqual(protection["INT/SC RatingBasis"], "0")
        self.assertEqual(protection["ConnectionType"], "2")
        self.assertEqual(protection["AF_ElectrodeConfig"], "0")
        self.assertEqual(protection["AF_EquipType"], "1")
        self.assertEqual(protection["AF_UnLinkedBoxDimensions"], "0")
        self.assertEqual(protection["AF_UnLinkedElectrodeConfig"], "0")
        self.assertEqual(protection["AF_UnLinkedEquipType"], "0")
        self.assertEqual(protection["ConnectedComponent1"], "FEED PNL L8C:2")
        self.assertEqual(protection["ConnectedComponent2"], "PNL L8C:1")
        self.assertEqual(protection["Frame/Rating"], "")
        self.assertEqual(utility["InService"], "1")
        self.assertEqual(utility["Phase"], "7")
        self.assertEqual(utility["SC Contribution Unit"], "2")
        self.assertEqual(utility["SC Contribution Type"], "0")
        self.assertEqual(survey_bus["AF_Inst/ZSI Avail"], "0")
        self.assertEqual(survey_bus["AF_ElectrodeConfig"], "1")
        self.assertEqual(survey_bus["AF_EquipType"], "1")
        self.assertEqual(survey_bus["AF_UnLinkedBoxDimensions"], "0")
        self.assertEqual(survey_bus["AF_UnLinkedElectrodeConfig"], "1")
        self.assertEqual(survey_bus["AF_UnLinkedEquipType"], "1")
        self.assertEqual(bus_duct["ContinuousRating"], "2000")
        self.assertEqual(bus_duct["Manufacturer"], "GE")
        self.assertEqual(bus_duct["Description"], "Spectra")

    def test_reports_unsupported_workbook_format(self) -> None:
        with self.assertRaisesRegex(ValueError, "requires an .xlsx or .xlsm"):
            build_records("survey.xls")


if __name__ == "__main__":
    unittest.main()